#!/usr/bin/env python3
"""
Lily Cogs + Online Inventory → Shopify (2 Oct 2026).

Sources (both in ~/Downloads unless given):
  Lily Cogs.xlsx                     product data + HA Online price (col K, EUR)
  Hartwick_StyleCosts_Aug2026.xlsx   "Online Iventory" tab only (S / M / L / XS units)

What it writes
  price      variant price (USD) = HA Online × EUR_USD, every variant of the product
  data       hartwick.category / core / dimensions / technique (read by the dev theme only)
  create     rows with no product → new DRAFT product (title = Old Name | Expression)
  stock      "available" at the store's location — needs read_inventory, write_inventory,
             read_locations on the CLI auth; skipped with a note until then

Matching: sheet (Old Name, Expression) ↔ product title's name part + hartwick.expression.
The Expression text is identical in the sheets and on the store, so no fuzzy matching.

    python3 scripts/costs-to-shopify.py --plan            # backup + plan, writes nothing
    python3 scripts/costs-to-shopify.py --write           # price + data + drafts (+ stock if scoped)
    python3 scripts/costs-to-shopify.py --stock --location "KAMLOOPS" [--zero-others]   # stock only
    python3 scripts/costs-to-shopify.py --revert FILE     # put prices back from a backup
"""
import argparse, csv, json, os, re, subprocess, sys, datetime, uuid
import openpyxl

STORE = "9c8a52-dc.myshopify.com"
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "data-import-2026-10-02")
DL = os.path.expanduser("~/Downloads")
EUR_USD = 1.14          # the workbook's own rate: USD RRP (col O) = EU RRP (col M) × 1.14

# ---------------------------------------------------------------------------
def gql(query, variables=None, mutate=False):
    args = ["shopify", "store", "execute", "-s", STORE, "--query", query]
    if variables is not None: args += ["--variables", json.dumps(variables)]
    if mutate: args.append("--allow-mutations")
    r = subprocess.run(args, capture_output=True, text=True)
    out = r.stdout if "{" in r.stdout else r.stdout + r.stderr
    i = out.find("{")
    try: d = json.loads(out[i:]) if i >= 0 else {}
    except json.JSONDecodeError: d = {"raw": out}
    return d.get("data", d)

def norm(x): return re.sub(r"[^a-z0-9]+", " ", str(x or "").lower()).strip()
def name_part(title): return norm(title.split("|")[0])
ALIAS = {"samarkant dress": "samarkand dress", "cedar trousers": "cedar pants", "sitges tote": "stiges tote bag",
         "cyprus pyjama set": "cyprus lounge set", "blossom shorts": "blossom lounge shorts",
         "jasmine dress": "jasmine lounge dress", "pine shirt": "pine lounge shirt",
         "jaipur scarf": "jaipur sarong", "calcutta blouse": "calcutta shirt",
         "classic stole 01": "stole 01", "classic stole 02": "stole 02",
         "dupatta 01": "oversized dupatta 01", "dupatta 02": "oversized dupatta 02"}

# ---------------------------------------------------------------------------
def load_lily(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    rows = []
    for r in ws.iter_rows(min_row=4, values_only=True):
        if not r[2] or not r[5]: continue
        rows.append({"category": r[0], "code": r[1] or (rows[-1]["code"] if rows and rows[-1]["name"] == r[2] else None),
                     "name": r[2], "dimensions": (str(r[3]).strip() if r[3] else None),
                     "core": str(r[4] or "").strip().lower() == "c", "expression": r[5], "technique": r[6],
                     "ha_online_eur": float(r[10]) if r[10] not in (None, "") else None})
    return rows

def load_inventory(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Online Iventory"]
    rows, name = [], None
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r[4] or r[2] == "total units": continue
        name = r[3] or r[2] or name
        rows.append({"name": name, "expression": r[4],
                     "qty": {k: (int(v) if isinstance(v, (int, float)) else None) for k, v in zip(["S", "M", "L", "XS"], r[6:10])}})
    return rows

PRODUCTS_Q = '''query($c:String){ products(first:50, after:$c){ pageInfo{hasNextPage endCursor} nodes{ id handle title status productType vendor tags
 style: metafield(namespace:"hartwick",key:"style"){ value }
 expr: metafield(namespace:"hartwick",key:"expression"){ value }
 options{ name values }
 variants(first:20){ nodes{ id title price compareAtPrice selectedOptions{name value} inventoryQuantity inventoryItem{ id } } } } } }'''

def pull_products():
    out, c = [], None
    while True:
        p = gql(PRODUCTS_Q, {"c": c})["products"]
        out += p["nodes"]
        if not p["pageInfo"]["hasNextPage"]: return out
        c = p["pageInfo"]["endCursor"]

def match(row, prods):
    nm = norm(row["name"]); nm = ALIAS.get(nm, nm); ex = norm(row["expression"])
    hits = [p for p in prods if p["expr"] and norm(p["expr"]["value"]) == ex and
            (name_part(p["title"]) == nm or name_part(p["title"]).replace("mayrlebone", "marylebone") == nm)]
    return hits

# stock: sheet column → store size option
def size_map(variants):
    sizes = [v["selectedOptions"][0]["value"] for v in variants]
    if sizes in (["Default Title"], ["One Size"]) or len(sizes) == 1: return {sizes[0]: "S"}
    if set(sizes) <= {"S/M", "M/L"}: return {"S/M": "S", "M/L": "M"}
    return {s: s for s in sizes}

# ---------------------------------------------------------------------------
def build(lily, inv, prods):
    plan, missing = [], []
    for r in lily:
        hits = match(r, prods)
        if not hits: missing.append(r); continue
        p = hits[0]
        usd = round(r["ha_online_eur"] * EUR_USD, 2) if r["ha_online_eur"] is not None else None
        plan.append({"row": r, "product": p, "dupes": hits[1:], "price_usd": usd})
    inv_plan = []
    for r in inv:
        hits = match(r, prods)
        if not hits: inv_plan.append({"row": r, "product": None}); continue
        p = hits[0]; sm = size_map(p["variants"]["nodes"]); changes, notes = [], []
        for v in p["variants"]["nodes"]:
            size = v["selectedOptions"][0]["value"]; col = sm.get(size)
            new = r["qty"].get(col) if col else None
            if new is None: notes.append(f"{size}: no count in sheet, left at {v['inventoryQuantity']}")
            else: changes.append({"variant": v, "size": size, "old": v["inventoryQuantity"], "new": new})
        used = set(sm.values())
        for col, q in r["qty"].items():
            if q is not None and col not in used: notes.append(f"sheet {col}={q} has no matching size on the product")
        inv_plan.append({"row": r, "product": p, "changes": changes, "notes": notes})
    return plan, missing, inv_plan

def write_reports(plan, missing, inv_plan, prods):
    os.makedirs(OUT, exist_ok=True)
    with open(os.path.join(OUT, "prices.csv"), "w", newline="") as f:
        w = csv.writer(f); w.writerow(["Style", "Name", "Expression", "Product", "Status", "HA Online EUR", "Old USD", "New USD", "Core", "Category", "Technique", "Dimensions"])
        for x in plan:
            r, p = x["row"], x["product"]
            w.writerow([r["code"], r["name"], r["expression"], p["title"], p["status"], r["ha_online_eur"],
                        p["variants"]["nodes"][0]["price"], x["price_usd"], "yes" if r["core"] else "", r["category"], r["technique"], r["dimensions"] or ""])
    with open(os.path.join(OUT, "stock.csv"), "w", newline="") as f:
        w = csv.writer(f); w.writerow(["Name", "Expression", "Product", "Size", "Old", "New", "Notes"])
        for x in inv_plan:
            r = x["row"]
            if not x["product"]: w.writerow([r["name"], r["expression"], "— no product —", "", "", "", json.dumps(r["qty"])]); continue
            for c in x["changes"]: w.writerow([r["name"], r["expression"], x["product"]["title"], c["size"], c["old"], c["new"], ""])
            for n in x["notes"]: w.writerow([r["name"], r["expression"], x["product"]["title"], "", "", "", n])

# ---------------------------------------------------------------------------
DEFS = [("category", "Category", "single_line_text_field", "Garment category from the costing sheet (Lily Cogs col A)."),
        ("core", "Core", "boolean", "Core range (Lily Cogs col E = c)."),
        ("dimensions", "Dimensions", "single_line_text_field", "Piece dimensions as given in Lily Cogs col D (silks, scarves)."),
        ("technique", "Technique", "single_line_text_field", "Technique family from Lily Cogs col G (Khadi, Ikat, Linen…).")]

def ensure_definitions():
    have = gql('{ metafieldDefinitions(first:100, ownerType:PRODUCT, namespace:"hartwick"){ nodes{ key } } }')
    keys = {n["key"] for n in (have.get("metafieldDefinitions") or {}).get("nodes", [])}
    q = "mutation($d:MetafieldDefinitionInput!){ metafieldDefinitionCreate(definition:$d){ createdDefinition{ key } userErrors{ message } } }"
    for k, name, t, desc in DEFS:
        if k in keys: continue
        r = gql(q, {"d": {"ownerType": "PRODUCT", "namespace": "hartwick", "key": k, "name": name, "type": t, "description": desc,
                          "access": {"storefront": "PUBLIC_READ"}}}, mutate=True)
        print(f"  definition hartwick.{k}: {r.get('metafieldDefinitionCreate')}")

def data_fields(gid, r):
    m = [{"ownerId": gid, "namespace": "hartwick", "key": "core", "type": "boolean", "value": "true" if r["core"] else "false"}]
    for k in ("category", "technique", "dimensions"):
        if r.get(k): m.append({"ownerId": gid, "namespace": "hartwick", "key": k, "type": "single_line_text_field", "value": str(r[k])})
    return m

def write_prices_and_data(plan):
    qp = "mutation($p:ID!,$v:[ProductVariantsBulkInput!]!){ productVariantsBulkUpdate(productId:$p, variants:$v){ userErrors{ field message } } }"
    qm = "mutation($m:[MetafieldsSetInput!]!){ metafieldsSet(metafields:$m){ userErrors{ field message } } }"
    for x in plan:
        p, r = x["product"], x["row"]
        msg = []
        if x["price_usd"] is not None:
            vs = [{"id": v["id"], "price": f'{x["price_usd"]:.2f}'} for v in p["variants"]["nodes"]]
            e = (gql(qp, {"p": p["id"], "v": vs}, mutate=True).get("productVariantsBulkUpdate") or {"userErrors": "no response"})["userErrors"]
            msg.append(f'price {p["variants"]["nodes"][0]["price"]} → {x["price_usd"]:.2f}' + (f" ✗ {e}" if e else ""))
        e = (gql(qm, {"m": data_fields(p["id"], r)}, mutate=True).get("metafieldsSet") or {"userErrors": "no response"})["userErrors"]
        msg.append("data ok" if not e else f"data ✗ {e}")
        print(f'  {p["title"][:60]:60} {"; ".join(msg)}')

def create_drafts(missing, prods, styles):
    q = "mutation($i:ProductSetInput!){ productSet(input:$i, synchronous:true){ product{ id title status } userErrors{ field message } } }"
    created = []
    for r in missing:
        sib = next((p for p in prods if name_part(p["title"]) == ALIAS.get(norm(r["name"]), norm(r["name"]))), None)
        opts = sib["options"] if sib else [{"name": "Title", "values": ["Default Title"]}]
        price = f'{round(r["ha_online_eur"] * EUR_USD, 2):.2f}' if r["ha_online_eur"] is not None else "0.00"
        o = opts[0]
        inp = {"title": f'{r["name"]} | {r["expression"]}', "status": "DRAFT", "vendor": "HARTWICK ATELIER",
               "productType": sib["productType"] if sib else "", "tags": (sib["tags"] if sib else []),
               "productOptions": [{"name": o["name"], "values": [{"name": v} for v in o["values"]]}],
               "variants": [{"optionValues": [{"optionName": o["name"], "name": v}], "price": price} for v in o["values"]],
               "metafields": [{"namespace": "hartwick", "key": "expression", "type": "single_line_text_field", "value": r["expression"]}]
                             + [{k: f[k] for k in ("namespace", "key", "type", "value")} for f in data_fields("x", r)]}
        st = (sib or {}).get("style", {}) or {}
        sid = st.get("value") or styles.get(norm(r["code"] or ""))
        if sid: inp["metafields"].append({"namespace": "hartwick", "key": "style", "type": "metaobject_reference", "value": sid})
        res = gql(q, {"i": inp}, mutate=True).get("productSet") or {}
        print(f'  + DRAFT {inp["title"]} @ {price} → {res.get("product") or res.get("userErrors") or res}')
        created.append({"input": inp, "result": res})
    return created

LEVELS_Q = """query($ids:[ID!]!){ nodes(ids:$ids){ ... on InventoryItem{ id
  inventoryLevels(first:20){ nodes{ location{ id name fulfillsOnlineOrders } quantities(names:["available"]){ quantity } } } } } }"""

def locations():
    d = gql('{ locations(first:20){ nodes{ id name isActive fulfillsOnlineOrders } } }')
    return [n for n in ((d.get("locations") or {}).get("nodes") or []) if n["isActive"]]

def write_stock(inv_plan, location=None, zero_others=False, dry=False):
    """Sheet count → `location`. With zero_others, every other online-selling
    location is set to 0 for the same items, so the storefront total = the sheet."""
    locs = locations()
    if not locs:
        print("  stock: SKIPPED — the CLI login cannot read locations (needs read_locations, read_inventory, write_inventory)."); return
    loc = next((l for l in locs if l["name"].lower() == (location or "").lower()), None)
    if not loc:
        if len(locs) == 1: loc = locs[0]
        else:
            print("  stock: SKIPPED — the store has %d active locations; say which one holds online stock:" % len(locs))
            for l in locs: print(f"     --location \"{l['name']}\"" + ("" if l["fulfillsOnlineOrders"] else "   (does not sell online)"))
            print("  add --zero-others to set every other online-selling location to 0 for these items."); return
    changes = [c for x in inv_plan if x["product"] for c in x["changes"]]
    ids = list({c["variant"]["inventoryItem"]["id"] for c in changes})
    levels = {}
    for i in range(0, len(ids), 50):
        for n in (gql(LEVELS_Q, {"ids": ids[i:i+50]}).get("nodes") or []):
            if n: levels[n["id"]] = {l["location"]["id"]: (l["location"], l["quantities"][0]["quantity"]) for l in n["inventoryLevels"]["nodes"]}
    sets, activate, rows = [], [], []
    for c in changes:
        iid = c["variant"]["inventoryItem"]["id"]; lv = levels.get(iid, {})
        cur = lv.get(loc["id"], (None, None))[1]
        if loc["id"] not in lv: activate.append(iid)
        if cur != c["new"]: sets.append({"inventoryItemId": iid, "locationId": loc["id"], "quantity": c["new"], "changeFromQuantity": None})
        rows.append([c["variant"]["id"], c["size"], loc["name"], cur, c["new"]])
        if zero_others:
            for lid, (l, q) in lv.items():
                if lid != loc["id"] and l["fulfillsOnlineOrders"] and q:
                    sets.append({"inventoryItemId": iid, "locationId": lid, "quantity": 0, "changeFromQuantity": None})
                    rows.append([c["variant"]["id"], c["size"], l["name"], q, 0])
    with open(os.path.join(OUT, "stock-by-location.csv"), "w", newline="") as f:
        w = csv.writer(f); w.writerow(["Variant", "Size", "Location", "Available now", "New"]); w.writerows(rows)
    print(f"  stock → {loc['name']}: {len(sets)} level changes, {len(activate)} items to stock there first"
          + (" (other online locations zeroed)" if zero_others else "") + "; see stock-by-location.csv")
    if dry: return
    # inventory mutations need @idempotent(key:) on current API versions; one fresh key per call
    qa = "mutation($i:ID!,$l:ID!,$k:String!){ inventoryActivate(inventoryItemId:$i, locationId:$l) @idempotent(key:$k){ userErrors{ message } } }"
    for iid in activate:
        r = gql(qa, {"i": iid, "l": loc["id"], "k": str(uuid.uuid4())}, mutate=True)
        if "inventoryActivate" not in r: r = gql(qa.replace(" @idempotent(key:$k)", "").replace(",$k:String!", ""), {"i": iid, "l": loc["id"]}, mutate=True)
        e = (r.get("inventoryActivate") or {"userErrors": str(r)[:300]}).get("userErrors")
        if e: print("  activate", iid, e)
    q = "mutation($i:InventorySetQuantitiesInput!,$k:String!){ inventorySetQuantities(input:$i) @idempotent(key:$k){ userErrors{ field message } } }"
    for i in range(0, len(sets), 100):
        r = gql(q, {"k": str(uuid.uuid4()), "i": {"name": "available", "reason": "correction", "referenceDocumentUri": "hartwick://online-inventory/2026-10-02",
                          "quantities": sets[i:i+100]}}, mutate=True)
        e = (r.get("inventorySetQuantities") or {}).get("userErrors")
        print(f"  stock {i+1}-{i+len(sets[i:i+100])}: " + ("ok" if r.get("inventorySetQuantities") and not e else str(e or r)[:400]))

def revert(path):
    snap = json.load(open(path))
    qp = "mutation($p:ID!,$v:[ProductVariantsBulkInput!]!){ productVariantsBulkUpdate(productId:$p, variants:$v){ userErrors{ message } } }"
    for p in snap:
        vs = [{"id": v["id"], "price": v["price"]} for v in p["variants"]["nodes"]]
        print(p["title"][:60], gql(qp, {"p": p["id"], "v": vs}, mutate=True).get("productVariantsBulkUpdate"))

# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--lily", default=os.path.join(DL, "Lily Cogs.xlsx"))
    ap.add_argument("--costs", default=os.path.join(DL, "Hartwick_StyleCosts_Aug2026.xlsx"))
    ap.add_argument("--plan", action="store_true"); ap.add_argument("--write", action="store_true")
    ap.add_argument("--stock", action="store_true"); ap.add_argument("--revert")
    ap.add_argument("--location", help='location that holds online stock, e.g. "KAMLOOPS"')
    ap.add_argument("--zero-others", action="store_true", help="set the other online-selling locations to 0 for these items")
    a = ap.parse_args()
    if a.revert: return revert(a.revert)
    lily, inv = load_lily(a.lily), load_inventory(a.costs)
    prods = pull_products()
    os.makedirs(OUT, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    json.dump(prods, open(os.path.join(OUT, f"backup-before-{stamp}.json"), "w"), indent=1)
    plan, missing, inv_plan = build(lily, inv, prods)
    write_reports(plan, missing, inv_plan, prods)
    print(f"Lily rows {len(lily)} · matched {len(plan)} · no product {len(missing)}: " + "; ".join(f'{r["name"]} / {r["expression"]}' for r in missing))
    print(f"Inventory rows {len(inv)} · matched {sum(1 for x in inv_plan if x['product'])}")
    for x in plan:
        if x["dupes"]: print("  duplicate on store:", [d["title"] + " (" + d["status"] + ")" for d in x["dupes"]])
    if a.write:
        ensure_definitions()
        write_prices_and_data(plan)
        mos = gql('{ metaobjects(type:"style", first:100){ nodes{ id code: field(key:"code"){ value } } } }')
        styles = {norm((m.get("code") or {}).get("value")): m["id"] for m in (mos.get("metaobjects") or {}).get("nodes", [])}
        create_drafts(missing, prods, styles)
    if a.write or a.stock or a.plan:
        write_stock(inv_plan, a.location, a.zero_others, dry=a.plan and not (a.write or a.stock))

if __name__ == "__main__":
    main()
